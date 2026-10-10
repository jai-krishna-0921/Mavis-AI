"""The studio renderers: build rich and edge-case specs, reopen the files and check their structure."""

from __future__ import annotations

import datetime
import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from mavis.studio.render import main, render
from mavis.studio.render_docx import render_doc
from mavis.studio.render_pptx import render_deck
from mavis.studio.render_xlsx import render_sheet
from mavis.studio.spec import (
    Block,
    Chart,
    Column,
    DeckSpec,
    DocSpec,
    Section,
    Series,
    SheetSpec,
    Slide,
    Stat,
    Step,
    Tab,
)
from mavis.studio.themes import THEMES, theme_for

THEME = theme_for("corporate", company="Acme Co")
CHART = Chart(kind="column", categories=["Q1", "Q2", "Q3"],
              series=[Series(name="Revenue", values=[10, 14, 21])], number_format="#,##0")


def _deck() -> DeckSpec:
    return DeckSpec(title="Growth plan", subtitle="FY27", slides=[
        Slide(layout="title", title="Growth plan", subtitle="Where we go next", notes="Open warmly"),
        Slide(layout="section", title="Where we are"),
        Slide(layout="bullets", title="Three things matter", bullets=["Retention", "Pricing", "Speed"]),
        Slide(layout="two_column", title="Before and after", left_heading="Before", left=["Manual"],
              right_heading="After", right=["Automated", "Faster"]),
        Slide(layout="stats", title="By the numbers",
              stats=[Stat(value="18 lakh", label="Revenue"), Stat(value="3x", label="Growth"),
                     Stat(value="92%", label="Retention")]),
        Slide(layout="quote", quote="Make it simple.", attribution="Founder"),
        Slide(layout="timeline", title="Roadmap",
              steps=[Step(when="Jan", what="Plan"), Step(when="Feb", what="Build"),
                     Step(when="Mar", what="Ship")]),
        Slide(layout="chart", title="Revenue", chart=CHART, takeaway="Q3 is the high point",
              notes="Explain Q3"),
        Slide(layout="chart", title="Mix", chart=CHART.model_copy(update={"kind": "doughnut"})),
        Slide(layout="chart", title="Trend", chart=CHART.model_copy(update={"kind": "line"})),
        Slide(layout="chart", title="Bars", chart=CHART.model_copy(update={"kind": "bar"})),
        Slide(layout="table", title="Plans", table_header=["Plan", "Price"],
              table_rows=[["Basic", "10"], ["Pro", "25"]], takeaway="Pro wins"),
        Slide(layout="closing", title="Thank you", subtitle="Questions?", bullets=["Email us"]),
    ])


def test_deck_structure(tmp_path: Path) -> None:
    out = render_deck(_deck(), THEME, tmp_path / "d.pptx", author="Jo")
    prs = Presentation(str(out))
    assert len(prs.slides) == 13
    assert prs.core_properties.title == "Growth plan"
    assert prs.core_properties.author == "Jo"
    assert (prs.slide_width, prs.slide_height) == (12192000, 6858000) or prs.slide_width > prs.slide_height
    assert prs.slides[0].notes_slide.notes_text_frame.text == "Open warmly"
    chart_slide = prs.slides[7]
    assert chart_slide.notes_slide.notes_text_frame.text == "Explain Q3"
    charts = [sh.chart for sh in chart_slide.shapes if sh.has_chart]
    assert len(charts) == 1
    assert list(charts[0].plots[0].categories) == ["Q1", "Q2", "Q3"]
    for idx in (8, 9, 10):
        assert any(sh.has_chart for sh in prs.slides[idx].shapes)
    table = next(sh.table for sh in prs.slides[11].shapes if sh.has_table)
    assert table.cell(0, 0).text == "Plan"
    assert table.cell(2, 1).text == "25"
    texts = [sh.text_frame.text for sh in prs.slides[2].shapes if sh.has_text_frame]
    assert "Retention" in texts
    assert "Acme Co" in texts


def test_deck_every_theme(tmp_path: Path) -> None:
    for name, theme in THEMES.items():
        render_deck(_deck(), theme, tmp_path / f"{name}.pptx")
        assert len(Presentation(str(tmp_path / f"{name}.pptx")).slides) == 13


def _doc() -> DocSpec:
    return DocSpec(title="Annual report", subtitle="FY27 review",
                   summary="Revenue grew.\n- Margin up\n- Churn down", sections=[
        Section(heading="Overview", blocks=[
            Block(kind="paragraph", text="First para.\n\nSecond para."),
            Block(kind="bullets", items=["One", "Two"]),
            Block(kind="numbered", items=["Alpha", "Beta"]),
            Block(kind="numbered", items=["Gamma"]),
            Block(kind="table", table_header=["Region", "Sales"],
                  table_rows=[["North", "10"], ["South", "7"]]),
            Block(kind="callout", text="Watch the margin."),
            Block(kind="quote", text="Less, but better."),
        ]),
        Section(heading="Detail", level=2, blocks=[Block(kind="paragraph", text="More.")]),
    ])


def test_doc_structure(tmp_path: Path) -> None:
    theme = theme_for("bold", company="Acme Co")
    out = render_doc(_doc(), theme, tmp_path / "d.docx", author="Jo", date_text="1 Jan 2027")
    doc = Document(str(out))
    assert doc.styles["Normal"].font.name == theme.body_font
    assert doc.styles["Heading 1"].font.name == theme.heading_font
    heads = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith(("Heading", "Title"))]
    assert ("Title", "Annual report") in heads
    assert ("Heading 1", "Overview") in heads
    assert ("Heading 2", "Detail") in heads
    # summary box, callout and data table
    assert len(doc.tables) == 3
    assert doc.tables[0].cell(0, 0).text.startswith("Key takeaways")
    data = next(t for t in doc.tables if t.cell(0, 0).text == "Region")
    assert data.cell(2, 1).text == "7"
    assert 'w:fill="' + theme.primary.lstrip("#") + '"' in data.cell(0, 0)._tc.xml
    styles = {p.style.name for p in doc.paragraphs}
    assert {"List Bullet", "List Number"} <= styles
    nums = {p._p.pPr.numPr.numId.val for p in doc.paragraphs if p.style.name == "List Number"}
    assert len(nums) == 2  # each numbered list restarts
    assert "PAGE" in doc.sections[0].footer._element.xml
    assert "Acme Co" in doc.sections[0].footer.paragraphs[0].text
    assert doc.core_properties.title == "Annual report"
    assert any("Jo" in p.text and "1 Jan 2027" in p.text for p in doc.paragraphs)


def _sheet() -> SheetSpec:
    cols = [Column(name="Item"), Column(name="Qty", kind="number", total=True),
            Column(name="Price", kind="currency", currency="Rs", total=True),
            Column(name="Margin", kind="percent"), Column(name="Due", kind="date")]
    rows = [["Pen", 3, 10.5, "12%", "2027-01-05"], ["Book", "1,200", "99", 0.2, "2027-02-01"],
            ["Bag", 2, 5, None, "not a date"]]
    chart = Chart(kind="column", categories=["Pen", "Book", "Bag"],
                  series=[Series(name="Qty", values=[3, 1200, 2])])
    free = Chart(kind="pie", categories=["A", "B"], series=[Series(name="Share", values=[1, 2])])
    return SheetSpec(title="Stock", tabs=[
        Tab(name="Stock: [2027]/Q1", columns=cols, rows=rows, chart=chart),
        Tab(name="stock: [2027]/q1", columns=cols[:2], rows=[["x", 1]], chart=free),
        Tab(name="A" * 31, columns=cols[:1]),
    ])


def test_sheet_structure(tmp_path: Path) -> None:
    theme = theme_for("minimal")
    out = render_sheet(_sheet(), theme, tmp_path / "s.xlsx")
    wb = load_workbook(str(out))
    assert wb.properties.title == "Stock"
    assert len(wb.sheetnames) == 3
    assert len({n.lower() for n in wb.sheetnames}) == 3
    assert all(len(n) <= 31 and not set("[]:*?/\\") & set(n) for n in wb.sheetnames)
    ws = wb.worksheets[0]
    head = ws["A1"]
    assert head.font.bold and head.fill.start_color.rgb.endswith(theme.primary.lstrip("#"))
    assert head.alignment.horizontal == "center"
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:E4"
    assert ws["B2"].number_format == "#,##0.##" and ws["B3"].value == 1200
    assert ws["C2"].number_format == '"Rs"#,##0.00'
    assert ws["D2"].number_format == "0.0%" and abs(ws["D2"].value - 0.12) < 1e-9
    assert ws["E2"].number_format == "yyyy-mm-dd" and ws["E2"].value.date() == datetime.date(2027, 1, 5)
    assert ws["E4"].value == "not a date"
    assert ws["A5"].value == "Total"
    assert ws["B5"].value == "=SUM(B2:B4)" and ws["C5"].value == "=SUM(C2:C4)"
    assert ws["B5"].font.bold and ws["B5"].border.top.style
    assert ws["D5"].value is None
    assert len(ws._charts) == 1
    assert len(wb.worksheets[1]._charts) == 1  # chart data written beside it when not a table column


def test_robustness(tmp_path: Path) -> None:
    long = "word " * 80
    deck = DeckSpec(title="T" * 120, slides=[
        Slide(layout="title"), Slide(layout="section"), Slide(layout="bullets"),
        Slide(layout="bullets", title=long[:120], bullets=[long[:400]] * 6),
        Slide(layout="two_column"), Slide(layout="stats"), Slide(layout="quote"), Slide(layout="timeline"),
        Slide(layout="chart", title="No chart", bullets=["Fallback"], takeaway="Only a takeaway"),
        Slide(layout="table", title="Ragged", table_header=["A"], table_rows=[["1", "2", "3"], [], ["x"]]),
        Slide(layout="table"), Slide(layout="closing"),
        Slide(layout="stats", stats=[Stat(value="1234567890123", label=long[:120])] * 4),
        Slide(layout="timeline", steps=[Step(when="w", what=long[:400])] * 6),
    ])
    assert len(Presentation(str(render_deck(deck, THEME, tmp_path / "e.pptx"))).slides) == 14
    doc = DocSpec(title="Edge", cover=False, sections=[
        Section(heading="", blocks=[Block(kind="table"), Block(kind="bullets"), Block(kind="callout"),
                                    Block(kind="quote"), Block(kind="paragraph", text=long * 10),
                                    Block(kind="table", table_header=["A"], table_rows=[["1", "2"], []])])])
    render_doc(doc, THEME, tmp_path / "e.docx")
    sheet = SheetSpec(title="Edge", tabs=[
        Tab(name="", columns=[Column(name="N", kind="number", total=True), Column(name="D", kind="date")],
            rows=[[None, None], ["abc", 5], [1, "2027-01-01"], [1]],
            chart=Chart(kind="line", categories=["a"], series=[Series(name="Nope", values=[1])])),
        Tab(name="Empty", columns=[Column(name="X", kind="currency", total=True)],
            chart=Chart(kind="doughnut", categories=["a"], series=[Series(name="X", values=[1])])),
    ])
    wb = load_workbook(str(render_sheet(sheet, THEME, tmp_path / "e.xlsx")))
    assert wb.sheetnames[0] == "Sheet"


def test_dispatch_and_cli(tmp_path: Path, capsys) -> None:
    out = render("deck", _deck(), THEME, tmp_path / "x.pptx")
    assert out.exists()
    spec_path = tmp_path / "spec.json"
    spec_path.write_text(json.dumps(json.loads(_doc().model_dump_json())))
    target = tmp_path / "cli.docx"
    assert main("doc", [str(spec_path), str(target), "--theme", "warm"]) == 0
    assert capsys.readouterr().out.strip() == str(target) and target.exists()
    assert main("doc", [str(tmp_path / "missing.json"), str(target)]) == 1
    assert main("doc", [str(spec_path)]) == 1


def test_column_text_shrinks_to_fit_its_count_and_wrapping():
    """Evals 2026-10-10: six two-line bullets overflowed a card at the size one bullet fits."""
    from mavis.studio.render_pptx import _fit_block

    one = ["Free $0, unlimited text, ads in US and most markets"]
    six = one * 6
    assert _fit_block(one, 5.3, 3.4, 20, 11, after_pt=10) == 20
    small = _fit_block(six, 5.3, 3.4, 20, 11, after_pt=10)
    assert 11 <= small < 20
    per_line = int(5.3 * 72 / (small * 0.52))
    lines = sum(-(-len(t) // per_line) for t in six)
    assert (lines * small * 1.2 + 50) / 72 <= 3.4


def test_table_cells_are_never_cut_and_rows_fit_the_slide():
    """Evals 2026-10-10: cells were clipped to 60 characters with an ellipsis at a fixed row height."""
    from mavis.studio.render_pptx import _table_rows

    long = "Typed notes, databases, templates; no handwriting or PDF annotation, AI meeting notes"
    grid = [["App", "Lecture", "Exam"]] + [["Notion", long, long]] * 3
    size, heights = _table_rows(grid, 3.68, 14, 3.95)
    assert 10 <= size <= 14 and sum(heights) <= 3.95
    assert heights[1] > heights[0]  # a wrapped row is taller than the one-line header
