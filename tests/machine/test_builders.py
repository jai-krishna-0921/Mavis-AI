"""Builders run inside a session (LocalSandbox here) from JSON data; never from formatted code."""

from __future__ import annotations

import io

import pytest

from mavis.domain.plans import DeckOutline, DocOutline, DocSection, SlideSpec
from mavis.machine.fake import MemoryWorkspaceStore
from mavis.machine.local import LocalSandbox
from mavis.machine.runtime import MachineRuntime
from mavis.store.repo import tasks


@pytest.fixture
async def rt(db, user, tmp_path, monkeypatch):
    from mavis.machine import quota

    monkeypatch.setattr(quota, "get_redis", lambda: None)

    async def deliver(*a, **k):
        return True

    runtime = MachineRuntime(LocalSandbox(root=tmp_path / "sb"), MemoryWorkspaceStore(), deliver=deliver)
    tid = await tasks.create(user.id, goal="make documents")
    return runtime, user.id, tid


async def _out(runtime, uid, tid, name):
    return await runtime.store.get(uid, f"out/{name}")


async def test_pptx_has_title_plus_slides_and_survives_odd_text(rt):
    from pptx import Presentation

    runtime, uid, tid = rt
    deck = DeckOutline(
        title="Standing desks 🚀",
        subtitle="Why and how",
        slides=[SlideSpec(title=f"Point {i}", bullets=["x" * 400, "", "ok ✅"]) for i in range(5)],
    )
    res = await runtime.build(uid, tid, "pptx", deck.model_dump(), "desks.pptx")
    assert res.ok, res.stderr
    prs = Presentation(io.BytesIO(await _out(runtime, uid, tid, "desks.pptx")))
    assert len(prs.slides) == 6


@pytest.mark.parametrize(
    "builder,name,check",
    [
        ("docx", "notes.docx", lambda b: b[:2] == b"PK"),
        ("pdf", "notes.pdf", lambda b: b[:4] == b"%PDF"),
    ],
)
async def test_doc_builders(rt, builder, name, check):
    runtime, uid, tid = rt
    outline = DocOutline(
        title="Trip notes",
        sections=[
            DocSection(
                heading="Day 1",
                paragraphs=["Café ☕ visit"],
                bullets=["a", "b"],
                table=[["k", "v"], ["1", "2"]],
            )
        ],
    )
    res = await runtime.build(uid, tid, builder, outline.model_dump(), name)
    assert res.ok, res.stderr
    assert check(await _out(runtime, uid, tid, name))


async def test_xlsx_and_chart(rt):
    from openpyxl import load_workbook

    runtime, uid, tid = rt
    sheets = {"sheets": [{"name": "Desks", "rows": [["Name", "Price"], ["A", 12999], ["B", 14500]]}]}
    assert (await runtime.build(uid, tid, "xlsx", sheets, "desks.xlsx")).ok
    wb = load_workbook(io.BytesIO(await _out(runtime, uid, tid, "desks.xlsx")))
    assert wb["Desks"]["B3"].value == 14500
    chart = {
        "kind": "line",
        "title": "Revenue",
        "x": ["Jan", "Feb", "Mar"],
        "series": [{"name": "north", "values": [1, 2, 3]}, {"name": "south", "values": [2, 2, 1]}],
    }
    assert (await runtime.build(uid, tid, "chart", chart, "revenue.png")).ok
    assert (await _out(runtime, uid, tid, "revenue.png"))[:4] == b"\x89PNG"


async def test_data_is_never_formatted_into_code(rt):
    runtime, uid, tid = rt
    evil = {"sheets": [{"name": "x'); import os; os.system('touch out/pwned'); ('", "rows": [["a"]]}]}
    await runtime.build(uid, tid, "xlsx", evil, "e.xlsx")
    assert await runtime.store.meta(uid, "out/pwned") is None


@pytest.mark.parametrize(
    "name,content,needle",
    [("a.csv", b"city,temp\nPune,31\n", "Pune"), ("b.txt", "Grüße aus Köln".encode(), "Köln")],
)
async def test_extract_text(rt, name, content, needle):
    from mavis.machine.ports import Provenance

    runtime, uid, tid = rt
    await runtime.store.put(uid, f"inbox/{name}", content, provenance=Provenance.USER_UPLOAD, cls=None)
    assert needle in await runtime.extract_text(uid, tid, f"inbox/{name}", 2000)


async def test_what_i_ran_lists_attempts_and_tail(rt):
    from mavis.machine.ports import ExecRequest

    runtime, uid, tid = rt
    await runtime.exec(uid, tid, ExecRequest(language="python", code="raise SystemExit(1)", timeout_s=20))
    await runtime.exec(
        uid, tid, ExecRequest(language="python", code="print('\\n'.join(map(str, range(40))))", timeout_s=20)
    )
    block = runtime.what_i_ran(tid)
    assert block.startswith("What I ran:")
    assert "attempt 1: exit 1" in block and "attempt 2: exit 0" in block
    assert "39" in block and "\n24\n" not in block  # only the last 15 lines
    assert "\u2014" not in block
