import json
import sys

from docx import Document

spec = json.load(open(sys.argv[1], encoding="utf-8"))
doc = Document()
doc.add_heading(str(spec.get("title", ""))[:200], level=0)
for sec in spec.get("sections", []):
    doc.add_heading(str(sec.get("heading", ""))[:200], level=1)
    for p in sec.get("paragraphs", []):
        doc.add_paragraph(str(p))
    for b in sec.get("bullets", []):
        doc.add_paragraph(str(b), style="List Bullet")
    table = sec.get("table") or []
    if table:
        t = doc.add_table(rows=len(table), cols=max(len(r) for r in table))
        t.style = "Table Grid"
        for i, row in enumerate(table):
            for j, v in enumerate(row):
                t.cell(i, j).text = str(v)
doc.save(sys.argv[2])
