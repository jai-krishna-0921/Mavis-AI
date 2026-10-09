import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

spec = json.load(open(sys.argv[1], encoding="utf-8"))
wb = Workbook()
wb.remove(wb.active)
for i, sheet in enumerate(spec.get("sheets", []) or [{"name": "Sheet1", "rows": []}]):
    name = (
        "".join(c for c in str(sheet.get("name") or f"Sheet{i + 1}") if c not in "[]:*?/\\")[:31]
        or f"Sheet{i + 1}"
    )
    ws = wb.create_sheet(name)
    for row in sheet.get("rows", []):
        ws.append([v if isinstance(v, int | float) else str(v) for v in row])
    for cell in ws[1] if ws.max_row else []:
        cell.font = Font(bold=True)
    for col in ws.columns:
        width = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(60, max(8, width + 2))
wb.save(sys.argv[2])
