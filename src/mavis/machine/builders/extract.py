import csv
import json
import sys

path, limit = sys.argv[2], int(json.load(open(sys.argv[1]))["max_chars"])
low = path.lower()
if low.endswith(".pdf"):
    from pypdf import PdfReader

    text = "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)
elif low.endswith(".docx"):
    from docx import Document

    text = "\n".join(p.text for p in Document(path).paragraphs)
elif low.endswith((".xlsx", ".xlsm")):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    text = "\n".join(
        f"# {ws.title}\n"
        + "\n".join(
            ",".join("" if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True)
        )
        for ws in wb.worksheets
    )
elif low.endswith(".csv"):
    with open(path, newline="", encoding="utf-8", errors="replace") as fh:
        text = "\n".join(",".join(r) for r in csv.reader(fh))
else:
    with open(path, encoding="utf-8", errors="replace") as fh:
        text = fh.read()
sys.stdout.write(text[:limit])
