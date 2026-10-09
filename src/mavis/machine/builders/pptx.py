import json
import sys

from pptx import Presentation
from pptx.util import Pt

spec = json.load(open(sys.argv[1], encoding="utf-8"))
prs = Presentation()
title = prs.slides.add_slide(prs.slide_layouts[0])
title.shapes.title.text = str(spec.get("title", ""))[:120]
if len(title.placeholders) > 1:
    title.placeholders[1].text = str(spec.get("subtitle", ""))[:200]
for s in spec.get("slides", []):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = str(s.get("title", ""))[:120]
    bullets = [str(b)[:220] for b in s.get("bullets", []) if str(b).strip()]
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, b in enumerate(bullets[:8]):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = b
        p.font.size = Pt(18)
    notes = str(s.get("notes", ""))
    if s.get("visual_hint"):
        notes += f"\nVisual idea: {s['visual_hint']}"
    if len(bullets) > 8:
        notes += "\nMore points: " + "; ".join(bullets[8:])
    slide.notes_slide.notes_text_frame.text = notes.strip()
prs.save(sys.argv[2])
