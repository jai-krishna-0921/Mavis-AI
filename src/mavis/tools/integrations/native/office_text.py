# Copyright (c) 2026 Jai Krishna. All rights reserved.
# Proprietary. Use is governed by LICENSE in the repository root.
"""Text out of an uploaded Excel, Word, PowerPoint or PDF file, so a Drive file the user made outside Google
reads like a Sheet or a Doc does.

The bytes are the user's own file but still untrusted input: Office files are zips, so the archive is
checked before parsing (entry count and inflated size) to refuse a zip bomb, and the text is capped.
"""

from __future__ import annotations

import io
import zipfile

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLSM = "application/vnd.ms-excel.sheet.macroEnabled.12"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
PDF = "application/pdf"
OFFICE_MIMES = frozenset({XLSX, XLSM, DOCX, PPTX, PDF})

MAX_ENTRIES = 2_000
MAX_INFLATED = 60 * 1024 * 1024
MAX_CHARS = 200_000
MAX_ROWS = 5_000  # per sheet
MAX_PAGES = 200


class UnreadableFile(ValueError):
    """The file is damaged, encrypted or too large inside to open safely."""


def _check_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile:
        raise UnreadableFile("not a valid Office file") from None
    if len(infos) > MAX_ENTRIES or sum(i.file_size for i in infos) > MAX_INFLATED:
        raise UnreadableFile("the file is too large inside to open")


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: list[str] = []
    try:
        for ws in wb.worksheets:
            out.append(f"# Sheet: {ws.title}")
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= MAX_ROWS:
                    out.append(f"(more rows not shown, first {MAX_ROWS})")
                    break
                cells = ["" if v is None else str(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if cells:
                    out.append(",".join(cells))
    finally:
        wb.close()
    return "\n".join(out)


def _docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    out = [p.text for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        out.extend(",".join(c.text.strip() for c in row.cells) for row in t.rows)
    return "\n".join(out)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    out: list[str] = []
    for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
        out.append(f"# Slide {i}")
        out.extend(sh.text_frame.text for sh in slide.shapes
                   if sh.has_text_frame and sh.text_frame.text.strip())
    return "\n".join(out)


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UnreadableFile("the PDF is password protected")
        return "\n".join((p.extract_text() or "") for p in reader.pages[:MAX_PAGES])
    except PdfReadError:
        raise UnreadableFile("not a valid PDF") from None


def office_text(data: bytes, mime: str) -> str:
    """The file's text (sheets as CSV rows), capped at MAX_CHARS. Raises UnreadableFile."""
    if mime == PDF:
        text = _pdf(data)
    else:
        _check_zip(data)
        reader = {XLSX: _xlsx, XLSM: _xlsx, DOCX: _docx, PPTX: _pptx}.get(mime)
        if reader is None:
            raise UnreadableFile(f"{mime} is not a supported file type")
        try:
            text = reader(data)
        except UnreadableFile:
            raise
        except (KeyError, ValueError, zipfile.BadZipFile, OSError) as exc:
            raise UnreadableFile(f"could not open the file ({type(exc).__name__})") from None
    return text[:MAX_CHARS]
