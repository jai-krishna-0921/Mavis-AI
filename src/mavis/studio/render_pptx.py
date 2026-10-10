"""Draw a DeckSpec as a 16:9 PPTX: a blank layout and plain shapes, so it opens cleanly in Google Slides.

Every layout has its own composition on one grid (0.6in margins, fixed title position, quiet footer).
Text never auto-shrinks reliably, so sizes are chosen from the text length and long text is clipped.
"""

from __future__ import annotations

import re
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from mavis.studio.spec import Chart, DeckSpec, Slide
from mavis.studio.themes import Theme

SW, SH = 13.333, 7.5
M = 0.6
CW = SW - 2 * M
TOP = 1.95  # content starts here on every content slide

_CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}
_NUMERIC = re.compile(r"^[\s$€£₹+\-(]*[\d.,]+\s*[%)xXkKmMbB]*\s*$")


@dataclass
class _Ctx:
    theme: Theme
    deck_title: str
    number: int = 0  # 1-based slide number
    section: int = 0  # running section counter


# --- small helpers -----------------------------------------------------------------------------------


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#").upper())


def _mix(a: str, b: str, t: float) -> str:
    """Blend colour a toward b by t (0..1), as '#RRGGBB'."""
    ca, cb = (tuple(int(c.lstrip("#")[i : i + 2], 16) for i in (0, 2, 4)) for c in (a, b))
    return "#" + "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(ca, cb, strict=True))


def _clip(text: str, limit: int) -> str:
    text = " ".join((text or "").split())
    if len(text) <= limit:
        return text
    return text[: limit - 3].rsplit(" ", 1)[0].rstrip(",;:") + "..."


def _tier(length: int, tiers: list[tuple[int, int]], floor: int) -> int:
    """Font size for a text of this length: the first tier whose limit it fits, else the floor."""
    for limit, size in tiers:
        if length <= limit:
            return size
    return floor


def _fit(text: str, width_in: float, max_pt: int, min_pt: int) -> int:
    """Largest size (pt) at which a single line of text fits the width."""
    need = max(len(text), 1) * 0.58 / 72
    return max(min_pt, min(max_pt, int(width_in / need)))


def _fit_block(items: list[str], width_in: float, height_in: float, max_pt: int, min_pt: int,
               after_pt: int = 0) -> int:
    """Largest size (pt) at which these paragraphs, wrapped to the width, fit the height: the count and
    the wrapping matter, not only the longest line (six two-line bullets overflow at a size one fits)."""
    for size in range(max_pt, min_pt - 1, -1):
        per_line = max(1, int(width_in * 72 / (size * 0.52)))  # average glyph is about half an em
        lines = sum(max(1, -(-len(t) // per_line)) for t in items)
        if (lines * size * 1.2 + after_pt * max(len(items) - 1, 0)) / 72 <= height_in:
            return size
    return min_pt


def _palette(theme: Theme, n: int) -> list[str]:
    """n series colours: the theme palette, then lighter tints of it."""
    base = theme.palette()
    out = list(base)
    step = 1
    while len(out) < n:
        out.extend(_mix(c, "#FFFFFF", min(0.25 * step, 0.7)) for c in base)
        step += 1
    return out[:n]


def _rect(slide, x: float, y: float, w: float, h: float, fill: str, *, shape=MSO_SHAPE.RECTANGLE,
          radius: float = 0.0):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False
    if radius and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    return shp


def _bullet(p, color: str) -> None:
    """Real hanging-indent bullet (survives import into Slides, unlike a typed glyph)."""
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "342900")
    pPr.set("indent", "-342900")
    clr = OxmlElement("a:buClr")
    srgb = OxmlElement("a:srgbClr")
    srgb.set("val", color.lstrip("#"))
    clr.append(srgb)
    char = OxmlElement("a:buChar")
    char.set("char", "•")
    pPr.append(clr)
    pPr.append(char)


def _text(slide, x: float, y: float, w: float, h: float, paras: str | list[str], *, font: str, size: int,
          color: str, bold: bool = False, italic: bool = False, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, after: int = 0, bullet: str = "", spacing: float = 0.0):
    """A zero-margin text box; one paragraph per string."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    items = [paras] if isinstance(paras, str) else paras
    for i, line in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if after:
            p.space_after = Pt(after)
        if spacing:
            p.line_spacing = spacing
        run = p.add_run()
        run.text = line
        f = run.font
        f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
        f.color.rgb = _rgb(color)
        if bullet:
            _bullet(p, bullet)
    return box


def _background(slide, color: str) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(color)


# --- shared chrome -----------------------------------------------------------------------------------


def _title(ctx: _Ctx, slide, text: str) -> None:
    t = ctx.theme
    _rect(slide, M, 0.5, 0.6, 0.07, t.accent)
    text = _clip(text, 120)
    size = _tier(len(text), [(40, 36), (60, 32), (90, 28)], 24)
    _text(slide, M, 0.68, CW, 1.1, text, font=t.heading_font, size=size, color=t.text, bold=True)


def _footer(ctx: _Ctx, slide) -> None:
    t = ctx.theme
    label = _clip(t.company or ctx.deck_title, 60)
    _rect(slide, M, 6.95, CW, 0.01, _mix(t.surface, t.muted, 0.25))
    _text(slide, M, 7.03, CW - 1.0, 0.25, label, font=t.body_font, size=10, color=t.muted)
    _text(slide, SW - M - 1.0, 7.03, 1.0, 0.25, str(ctx.number), font=t.body_font, size=10, color=t.muted,
          align=PP_ALIGN.RIGHT)


def _light(t: Theme) -> str:
    """Muted text colour that reads on the dark background."""
    return _mix(t.dark, "#FFFFFF", 0.72)


# --- layouts -----------------------------------------------------------------------------------------


def _title_slide(ctx: _Ctx, slide, s: Slide, deck: DeckSpec) -> None:
    t = ctx.theme
    _background(slide, t.dark)
    if t.company:
        _text(slide, M, 0.6, CW, 0.3, t.company.upper(), font=t.body_font, size=12, color=_light(t),
              bold=True)
    title = _clip(s.title or deck.title, 120)  # the spec's own limit: shrink, never cut a title short
    sub = _clip(s.subtitle or deck.subtitle, 200)
    size = _tier(len(title), [(24, 60), (46, 48), (76, 38)], 30)
    _rect(slide, M, 2.35, 1.1, 0.09, t.accent)
    _text(slide, M, 2.7, CW, 2.3, title, font=t.heading_font, size=size, color="#FFFFFF", bold=True,
          spacing=1.0)
    if sub:
        _text(slide, M, 5.15, 10.5, 1.2, sub, font=t.body_font, size=22 if len(sub) < 110 else 18,
              color=_light(t))


def _section_slide(ctx: _Ctx, slide, s: Slide) -> None:
    t = ctx.theme
    ctx.section += 1
    _background(slide, t.dark)
    _rect(slide, 0, 0, 0.35, SH, t.accent)
    _text(slide, 1.1, 1.7, 6, 1.6, f"{ctx.section:02d}", font=t.heading_font, size=110, color=t.accent,
          bold=True)
    title = _clip(s.title, 120)
    _text(slide, 1.1, 3.6, 10.8, 1.7, title, font=t.heading_font,
          size=_tier(len(title), [(30, 52), (55, 44), (80, 36)], 30), color="#FFFFFF", bold=True,
          anchor=MSO_ANCHOR.TOP)
    if s.subtitle:
        _text(slide, 1.1, 5.4, 10.2, 1.0, _clip(s.subtitle, 200), font=t.body_font, size=20, color=_light(t))


def _bullets_slide(ctx: _Ctx, slide, s: Slide, items: list[str] | None = None) -> None:
    t = ctx.theme
    _title(ctx, slide, s.title)
    rows = items if items is not None else [_clip(b, 200) for b in s.bullets if b.strip()]
    rows = rows[:6] or [_clip(x, 200) for x in (s.subtitle, s.takeaway) if x.strip()][:1]
    if not rows:
        return
    longest = max(len(r) for r in rows)
    size = _tier(longest, [(70, 24), (120, 22), (170, 20)], 18)
    row_h = min(0.95, 4.7 / len(rows))
    for i, row in enumerate(rows):
        y = TOP + i * row_h
        d = 0.46
        dot = _rect(slide, M, y + (row_h - d) / 2, d, d, t.primary, shape=MSO_SHAPE.OVAL)
        tf = dot.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(i + 1)
        r.font.name, r.font.size, r.font.bold = t.heading_font, Pt(16), True
        r.font.color.rgb = _rgb("#FFFFFF")
        _text(slide, M + 0.8, y, CW - 0.8, row_h, row, font=t.body_font, size=size, color=t.text,
              anchor=MSO_ANCHOR.MIDDLE)


def _two_column_slide(ctx: _Ctx, slide, s: Slide) -> None:
    t = ctx.theme
    _title(ctx, slide, s.title)
    gap = 0.4
    w = (CW - gap) / 2
    cols = [(s.left_heading, s.left, t.primary), (s.right_heading, s.right, t.secondary)]
    longest = max((len(b) for _, items, _ in cols for b in items), default=0)
    bodies = [[_clip(b, 160) for b in items if b.strip()][:6] for _, items, _ in cols]
    size = min([_tier(longest, [(60, 20), (100, 18)], 16)] +
               [_fit_block(b, w - 0.7, 3.4, 20, 11, after_pt=10) for b in bodies if b])
    for i, ((heading, _, color), body) in enumerate(zip(cols, bodies, strict=True)):
        x = M + i * (w + gap)
        _rect(slide, x, TOP, w, 4.7, t.surface)
        _rect(slide, x, TOP, w, 0.08, color)
        _text(slide, x + 0.35, TOP + 0.35, w - 0.7, 0.6, _clip(heading, 60), font=t.heading_font, size=22,
              color=color if color != t.secondary else t.text, bold=True)
        if body:
            _text(slide, x + 0.35, TOP + 1.15, w - 0.7, 3.4, body, font=t.body_font, size=size, color=t.text,
                  after=10, bullet=color)


def _stats_slide(ctx: _Ctx, slide, s: Slide) -> None:
    t = ctx.theme
    stats = [x for x in s.stats if x.value.strip() or x.label.strip()][:4]
    if not stats:
        return _bullets_slide(ctx, slide, s)
    _title(ctx, slide, s.title)
    n = len(stats)
    gap = 0.35
    w = (CW - gap * (n - 1)) / n
    for i, st in enumerate(stats):
        x = M + i * (w + gap)
        _rect(slide, x, 2.2, w, 3.6, t.surface)
        _rect(slide, x, 2.2, w, 0.08, t.primary)
        value = _clip(st.value, 14)
        _text(slide, x + 0.2, 2.8, w - 0.4, 1.4, value, font=t.heading_font,
              size=_fit(value, w - 0.4, 72 if n <= 3 else 60, 28), color=t.primary, bold=True,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        label = _clip(st.label, 70)
        _text(slide, x + 0.3, 4.4, w - 0.6, 1.2, label, font=t.body_font, size=18 if len(label) < 36 else 15,
              color=t.muted, align=PP_ALIGN.CENTER)
    note = _clip(s.takeaway or s.subtitle, 160)
    if note:
        _text(slide, M, 6.2, CW, 0.5, note, font=t.body_font, size=16, color=t.muted, align=PP_ALIGN.CENTER)


def _quote_slide(ctx: _Ctx, slide, s: Slide) -> None:
    t = ctx.theme
    quote = _clip(s.quote or s.subtitle or s.title, 300)
    _text(slide, M + 0.2, 0.75, 2.5, 2.0, "“", font=t.heading_font, size=160, color=t.accent, bold=True)
    size = _tier(len(quote), [(90, 32), (150, 28), (220, 24)], 20)
    _text(slide, 1.5, 2.55, SW - 3.0, 3.0, quote, font=t.heading_font, size=size, color=t.text,
          spacing=1.1)
    who = _clip(s.attribution, 80)
    if who:
        _rect(slide, 1.5, 5.85, 0.6, 0.06, t.accent)
        _text(slide, 1.5, 6.05, SW - 3.0, 0.4, who, font=t.body_font, size=18, color=t.muted)


def _timeline_slide(ctx: _Ctx, slide, s: Slide) -> None:
    t = ctx.theme
    steps = [x for x in s.steps if x.when.strip() or x.what.strip()][:6]
    if not steps:
        return _bullets_slide(ctx, slide, s)
    _title(ctx, slide, s.title)
    line_y, d = 4.35, 0.4
    _rect(slide, M, line_y - 0.02, CW, 0.04, _mix(t.surface, t.muted, 0.35))
    n = len(steps)
    col = CW / n
    bw = 5.0 if n == 1 else min(col * 1.85, 3.6)
    for i, st in enumerate(steps):
        cx = M + col * (i + 0.5)
        dot = _rect(slide, cx - d / 2, line_y - d / 2, d, d, t.primary, shape=MSO_SHAPE.OVAL)
        dot.line.color.rgb = _rgb(t.background)
        dot.line.width = Pt(3)
        tf = dot.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(i + 1)
        r.font.name, r.font.size, r.font.bold = t.heading_font, Pt(12), True
        r.font.color.rgb = _rgb("#FFFFFF")
        x = min(max(cx - bw / 2, M), SW - M - bw)
        above = i % 2 == 0
        y, anchor = (line_y - 0.45 - 1.85, MSO_ANCHOR.BOTTOM) if above else (line_y + 0.45, MSO_ANCHOR.TOP)
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(bw), Inches(1.85))
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        for j, (txt, size, bold, color, face) in enumerate([
            (_clip(st.when, 40), 18, True, t.primary, t.heading_font),
            (_clip(st.what, 130), 15, False, t.text, t.body_font),
        ]):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            p.space_after = Pt(4)
            r = p.add_run()
            r.text = txt
            r.font.name, r.font.size, r.font.bold = face, Pt(size), bold
            r.font.color.rgb = _rgb(color)


def _style_chart(ctx: _Ctx, chart_obj, spec: Chart, n_cats: int) -> None:
    t = ctx.theme
    kind = spec.kind
    pie = kind in ("pie", "doughnut")
    grid = _mix(t.background, t.muted, 0.25)
    chart_obj.has_title = False
    chart_obj.font.size = Pt(12)
    chart_obj.font.name = t.body_font
    chart_obj.font.color.rgb = _rgb(t.muted)
    chart_obj.has_legend = pie or len(spec.series) > 1
    if chart_obj.has_legend:
        chart_obj.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart_obj.legend.include_in_layout = False
        chart_obj.legend.font.size = Pt(12)
    plot = chart_obj.plots[0]
    colors = _palette(t, n_cats if pie else len(spec.series))
    labels = pie or n_cats * len(spec.series) <= 12
    if pie:
        for idx, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = _rgb(colors[idx])
            point.format.line.color.rgb = _rgb(t.background)
    else:
        for ser, color in zip(plot.series, colors, strict=False):
            if kind == "line":
                ser.format.line.color.rgb = _rgb(color)
                ser.format.line.width = Pt(3)
                ser.smooth = False
                ser.marker.style = XL_MARKER_STYLE.CIRCLE
                ser.marker.size = 8
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = _rgb(color)
                ser.marker.format.line.color.rgb = _rgb(color)
            else:
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = _rgb(color)
        if kind in ("bar", "column"):
            plot.gap_width = 60
            plot.overlap = -10 if len(spec.series) > 1 else 0
        cat = chart_obj.category_axis
        cat.has_major_gridlines = False
        cat.format.line.color.rgb = _rgb(grid)
        cat.tick_labels.font.size = Pt(12)
        if kind == "bar":
            cat.reverse_order = True
        val = chart_obj.value_axis
        val.format.line.fill.background()
        if labels:
            val.visible = False
            val.has_major_gridlines = False
        else:
            val.has_major_gridlines = True
            val.major_gridlines.format.line.color.rgb = _rgb(grid)
            val.major_gridlines.format.line.width = Pt(0.75)
            val.tick_labels.font.size = Pt(12)
            if spec.number_format:
                val.tick_labels.number_format = spec.number_format
                val.tick_labels.number_format_is_linked = False
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = Pt(13 if not pie else 14)
        dl.font.bold = True
        dl.font.color.rgb = _rgb("#FFFFFF" if kind == "doughnut" else t.text)
        if pie:
            dl.show_percentage = spec.number_format == ""
            dl.show_value = spec.number_format != ""
            dl.number_format = spec.number_format or "0%"
        else:
            dl.show_value = True
            dl.number_format = spec.number_format or "General"
        dl.number_format_is_linked = False
        if kind != "doughnut":
            dl.position = {"line": XL_LABEL_POSITION.ABOVE}.get(kind, XL_LABEL_POSITION.OUTSIDE_END)


def _chart_slide(ctx: _Ctx, slide, s: Slide) -> None:
    t = ctx.theme
    spec = s.chart
    if spec is None or not spec.categories or not spec.series:
        return _bullets_slide(ctx, slide, s)
    _title(ctx, slide, s.title)
    cats = [_clip(c, 30) for c in spec.categories]
    data = CategoryChartData(number_format=spec.number_format or "General")
    data.categories = cats
    series = spec.series[:1] if spec.kind in ("pie", "doughnut") else spec.series
    for ser in series:
        vals = list(ser.values)[: len(cats)]
        vals += [None] * (len(cats) - len(vals))
        data.add_series(_clip(ser.name, 40) or "Series", vals)
    note = _clip(s.takeaway, 160)
    h = 4.1 if note else 4.75
    frame = slide.shapes.add_chart(_CHART_TYPES[spec.kind], Inches(M), Inches(TOP - 0.1), Inches(CW),
                                   Inches(h), data)
    _style_chart(ctx, frame.chart, spec.model_copy(update={"series": series}), len(cats))
    if note:
        _rect(slide, M, 6.2, 0.08, 0.5, t.accent)
        _text(slide, M + 0.3, 6.2, CW - 0.3, 0.5, note, font=t.body_font, size=18, color=t.text, bold=True,
              anchor=MSO_ANCHOR.MIDDLE)


def _table_slide(ctx: _Ctx, slide, s: Slide) -> None:
    t = ctx.theme
    header = [_clip(h, 40) for h in s.table_header]
    body = [[_clip(c, 60) for c in r] for r in s.table_rows[:8]]
    ncols = max([len(header), *(len(r) for r in body)], default=0)
    if ncols == 0:
        return _bullets_slide(ctx, slide, s)
    _title(ctx, slide, s.title)
    header += [""] * (ncols - len(header))
    body = [r + [""] * (ncols - len(r)) for r in body]
    has_header = any(header)
    grid = ([header] if has_header else []) + body
    numeric = [bool(body) and all(_NUMERIC.match(r[c]) or not r[c] for r in body) for c in range(ncols)]
    longest = max((len(c) for r in grid for c in r), default=0)
    size = 16 if ncols <= 4 and longest <= 40 else 14 if longest <= 40 else 12
    row_h = 0.6 if len(grid) <= 7 else 0.5
    shape = slide.shapes.add_table(len(grid), ncols, Inches(M), Inches(TOP), Inches(CW),
                                   Inches(row_h * len(grid)))
    table = shape.table
    table.horz_banding = False
    table.first_row = False
    for r_i, row in enumerate(grid):
        table.rows[r_i].height = Inches(row_h)
        is_head = has_header and r_i == 0
        band = r_i - (1 if has_header else 0)
        for c_i, text in enumerate(row):
            cell = table.cell(r_i, c_i)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(t.primary if is_head else t.surface if band % 2 == 0 else
                                            t.background)
            cell.margin_left = cell.margin_right = Inches(0.18)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if numeric[c_i] else PP_ALIGN.LEFT
            run = p.add_run()
            run.text = text
            run.font.name = t.heading_font if is_head else t.body_font
            run.font.size = Pt(size)
            run.font.bold = is_head
            run.font.color.rgb = _rgb("#FFFFFF" if is_head else t.text)
    note = _clip(s.takeaway, 160)
    if note:
        y = TOP + row_h * len(grid) + 0.3
        _rect(slide, M, y, 0.08, 0.5, t.accent)
        _text(slide, M + 0.3, y, CW - 0.3, 0.5, note, font=t.body_font, size=18, color=t.text, bold=True,
              anchor=MSO_ANCHOR.MIDDLE)


def _closing_slide(ctx: _Ctx, slide, s: Slide, deck: DeckSpec) -> None:
    t = ctx.theme
    _background(slide, t.dark)
    title = _clip(s.title or "Thank you", 120)
    _rect(slide, M, 1.55, 1.1, 0.09, t.accent)
    _text(slide, M, 1.9, 11.5, 1.6, title, font=t.heading_font,
          size=_tier(len(title), [(30, 60), (55, 48), (80, 38)], 32), color="#FFFFFF", bold=True)
    sub = _clip(s.subtitle, 200)
    if sub:
        _text(slide, M, 3.6, 10.5, 0.9, sub, font=t.body_font, size=22, color=_light(t))
    steps = [_clip(b, 120) for b in s.bullets if b.strip()][:4]
    if steps:
        _text(slide, M, 4.6, 11, 2.2, steps, font=t.body_font, size=20, color="#FFFFFF", after=8,
              bullet=t.accent)
    if t.company:
        _text(slide, M, 6.9, CW, 0.3, t.company, font=t.body_font, size=12, color=_light(t))


# --- entry point -------------------------------------------------------------------------------------


def render_deck(spec: DeckSpec, theme: Theme, path: Path, *, author: str = "") -> Path:
    """Write the deck to `path` (a .pptx) and return it."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    blank = prs.slide_layouts[6]
    ctx = _Ctx(theme=theme, deck_title=spec.title)
    simple: dict[str, Callable[[_Ctx, object, Slide], None]] = {
        "section": _section_slide,
        "bullets": _bullets_slide,
        "two_column": _two_column_slide,
        "stats": _stats_slide,
        "quote": _quote_slide,
        "timeline": _timeline_slide,
        "chart": _chart_slide,
        "table": _table_slide,
    }
    for i, s in enumerate(spec.slides, start=1):
        ctx.number = i
        slide = prs.slides.add_slide(blank)
        _background(slide, theme.background)
        if s.layout == "title":
            _title_slide(ctx, slide, s, spec)
        elif s.layout == "closing":
            _closing_slide(ctx, slide, s, spec)
        else:
            simple[s.layout](ctx, slide, s)
            if s.layout != "section":
                _footer(ctx, slide)
        if s.notes.strip():
            slide.notes_slide.notes_text_frame.text = s.notes
    props = prs.core_properties
    props.title = spec.title
    props.author = author or theme.company or "Mavis AI"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path
