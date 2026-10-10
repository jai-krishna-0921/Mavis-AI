"""Draw a SheetSpec as an XLSX: one styled table per tab, real number formats, totals and a native chart."""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from mavis.studio.spec import Chart, Column, SheetSpec, Tab
from mavis.studio.themes import Theme

_BAD_NAME = re.compile(r"[\[\]:*?/\\]")
_NUMBER = re.compile(r"[^\d.\-+eE]")
MAX_WIDTH = 50


def _hex(color: str) -> str:
    return color.lstrip("#").upper()


def _mix(a: str, b: str, t: float) -> str:
    ca, cb = (tuple(int(c.lstrip("#")[i : i + 2], 16) for i in (0, 2, 4)) for c in (a, b))
    return "#" + "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(ca, cb, strict=True))


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=_hex(color), end_color=_hex(color))


def _sheet_name(raw: str, used: set[str]) -> str:
    """A legal, unique (case-insensitive) worksheet name of at most 31 characters."""
    base = _BAD_NAME.sub(" ", raw).strip().strip("'").strip() or "Sheet"
    name, n = base[:31], 1
    while name.lower() in used:
        n += 1
        suffix = f" {n}"
        name = base[: 31 - len(suffix)] + suffix
    used.add(name.lower())
    return name


def _number_format(col: Column, whole: bool = False) -> str:
    """`whole`: every value in the column is an integer. "#,##0.##" would show 25000 as "25,000." in
    Excel, so a number column is either whole or two places."""
    if col.kind == "number":
        return "#,##0" if whole else "#,##0.00"
    if col.kind == "currency":
        return f'"{col.currency}"#,##0.00' if col.currency else "#,##0.00"
    if col.kind == "percent":
        return "0.0%"
    if col.kind == "date":
        return "yyyy-mm-dd"
    return "General"


def _to_float(v: str | float | int) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, int | float):
        return float(v)
    cleaned = _NUMBER.sub("", v.replace(",", ""))
    try:
        return float(cleaned)
    except ValueError:
        return None


def _coerce(value: str | float | int | None, col: Column):
    """The cell value for a column kind; anything unparseable stays as given."""
    if value is None or value == "":
        return None
    if col.kind in ("number", "currency"):
        num = _to_float(value)
        return value if num is None else num
    if col.kind == "percent":
        num = _to_float(value)
        if num is None:
            return value
        # "12%" and 12 mean twelve percent; 0.12 is already a fraction
        return num / 100 if (isinstance(value, str) and "%" in value) or abs(num) > 1 else num
    if col.kind == "date":
        if isinstance(value, str):
            text = value.strip()
            for parse in (date.fromisoformat, lambda x: datetime.fromisoformat(x).date(),
                          lambda x: date.fromisoformat(x[:10])):
                try:
                    return parse(text)
                except ValueError:
                    continue
        return value
    return value


def _display_len(value, fmt: str) -> int:
    if value is None:
        return 0
    if isinstance(value, datetime | date):
        return 10
    if isinstance(value, int | float):
        return len(f"{value:,.2f}") + (len(fmt.split("#")[0]) if fmt.startswith('"') else 0)
    return max((len(line) for line in str(value).splitlines()), default=0)


def _colors(theme: Theme, n: int) -> list[str]:
    base = theme.palette()
    out = list(base)
    step = 1
    while len(out) < n:
        out.extend(_mix(c, "#FFFFFF", min(0.25 * step, 0.7)) for c in base)
        step += 1
    return [_hex(c) for c in out[:n]]


def _build_chart(ws, tab: Tab, spec: Chart, theme: Theme, last_row: int, anchor_col: int) -> None:
    """Chart beside the table; it references the table when its series are table columns, else a data block
    written under the chart."""
    names = {c.name.strip().lower(): i + 1 for i, c in enumerate(tab.columns)}
    by_table = bool(tab.rows) and all(s.name.strip().lower() in names for s in spec.series)
    pie = spec.kind in ("pie", "doughnut")
    series = spec.series[:1] if pie else spec.series
    chart = {"bar": BarChart, "column": BarChart, "line": LineChart, "pie": PieChart,
             "doughnut": DoughnutChart}[spec.kind]()
    if isinstance(chart, BarChart):
        chart.type = "bar" if spec.kind == "bar" else "col"
        chart.gapWidth = 60
    chart.title = None
    chart.width, chart.height = 20, 10
    chart.legend.position = "b"
    if len(series) == 1 and not pie:
        chart.legend = None
    if by_table:
        cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
        for s in series:
            col = names[s.name.strip().lower()]
            chart.add_data(Reference(ws, min_col=col, min_row=1, max_row=last_row), titles_from_data=True)
        n_cats = last_row - 1
    else:
        top = 24  # below the chart, same columns
        muted = Font(bold=True, color=_hex(theme.muted))
        ws.cell(row=top, column=anchor_col, value="Chart data").font = muted
        for j, s in enumerate(series):
            ws.cell(row=top + 1, column=anchor_col + 1 + j, value=s.name).font = muted
        for i, cat in enumerate(spec.categories):
            ws.cell(row=top + 2 + i, column=anchor_col, value=cat)
            for j, s in enumerate(series):
                if i < len(s.values):
                    ws.cell(row=top + 2 + i, column=anchor_col + 1 + j, value=s.values[i])
        n_cats = len(spec.categories)
        end = top + 1 + n_cats
        cats = Reference(ws, min_col=anchor_col, min_row=top + 2, max_row=end)
        for j in range(len(series)):
            chart.add_data(Reference(ws, min_col=anchor_col + 1 + j, min_row=top + 1, max_row=end),
                           titles_from_data=True)
        ws.column_dimensions[get_column_letter(anchor_col)].width = 18
    chart.set_categories(cats)
    colors = _colors(theme, n_cats if pie else len(series))
    for i, s in enumerate(chart.series):
        if pie:
            for k in range(n_cats):
                pt = DataPoint(idx=k)
                pt.graphicalProperties.solidFill = colors[k]
                pt.graphicalProperties.line.solidFill = "FFFFFF"
                s.dPt.append(pt)
        elif spec.kind == "line":
            s.graphicalProperties.line.solidFill = colors[i]
            s.graphicalProperties.line.width = 28575
            s.marker.symbol = "circle"
            s.marker.graphicalProperties.solidFill = colors[i]
            s.marker.graphicalProperties.line.solidFill = colors[i]
            s.smooth = False
        else:
            s.graphicalProperties.solidFill = colors[i]
            s.graphicalProperties.line.solidFill = colors[i]
    if pie:
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        for attr in ("showVal", "showSerName", "showCatName", "showLegendKey"):
            setattr(chart.dataLabels, attr, False)
    else:
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        if spec.number_format:
            chart.y_axis.number_format = spec.number_format
        chart.y_axis.majorGridlines.spPr = None
    ws.add_chart(chart, f"{get_column_letter(anchor_col)}2")


def _write_tab(ws, tab: Tab, theme: Theme) -> None:
    cols = tab.columns
    head_font = Font(name=theme.body_font, bold=True, color="FFFFFF", size=11)
    body_font = Font(name=theme.body_font, color=_hex(theme.text), size=11)
    thin = Side(style="thin", color=_hex(_mix(theme.background, theme.muted, 0.3)))
    ws.sheet_properties.tabColor = _hex(theme.primary)
    ws.sheet_view.showGridLines = False
    widths = [0] * len(cols)

    for c_i, col in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c_i, value=col.name)
        cell.font, cell.fill = head_font, _fill(theme.primary)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        widths[c_i - 1] = len(col.name) + 4  # room for the filter button
    ws.row_dimensions[1].height = 22

    coerced = [[_coerce(row[c_i] if c_i < len(row) else None, col) for c_i, col in enumerate(cols)]
               for row in tab.rows]
    whole = [all(float(r[c_i]).is_integer() for r in coerced if isinstance(r[c_i], int | float))
             for c_i in range(len(cols))]
    for r_i, row in enumerate(coerced, start=2):
        for c_i, col in enumerate(cols, start=1):
            value = row[c_i - 1]
            fmt = _number_format(col, whole[c_i - 1])
            cell = ws.cell(row=r_i, column=c_i, value=value)
            cell.font = body_font
            if not isinstance(value, str):
                cell.number_format = fmt
            if r_i % 2 == 0:
                cell.fill = _fill(theme.surface)
            numeric = col.kind in ("number", "currency", "percent") and not isinstance(value, str)
            cell.alignment = Alignment(
                horizontal="right" if numeric else "center" if col.kind == "date" else "left",
                vertical="center", wrap_text=col.kind == "text")
            widths[c_i - 1] = max(widths[c_i - 1], _display_len(value, fmt) + 2)

    last = len(tab.rows) + 1
    if any(c.total for c in cols) and tab.rows:
        total_row = last + 1
        for c_i, col in enumerate(cols, start=1):
            cell = ws.cell(row=total_row, column=c_i)
            cell.font = Font(name=theme.body_font, bold=True, color=_hex(theme.text), size=11)
            cell.border = Border(top=Side(style="medium", color=_hex(theme.primary)))
            if col.total:
                letter = get_column_letter(c_i)
                cell.value = f"=SUM({letter}2:{letter}{last})"
                cell.number_format = _number_format(col, whole[c_i - 1])
                cell.alignment = Alignment(horizontal="right", vertical="center")
                widths[c_i - 1] = max(widths[c_i - 1], 14)
            elif c_i == 1:
                cell.value = "Total"
        if not cols[0].total:
            widths[0] = max(widths[0], 8)
    elif tab.rows:
        for c_i in range(1, len(cols) + 1):
            ws.cell(row=last, column=c_i).border = Border(bottom=thin)

    for c_i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c_i)].width = min(MAX_WIDTH, max(10, w + 1))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(last, 1)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    if tab.chart is not None and tab.chart.categories and tab.chart.series:
        _build_chart(ws, tab, tab.chart, theme, last, len(cols) + 2)


def render_sheet(spec: SheetSpec, theme: Theme, path: Path) -> Path:
    """Write the workbook to `path` (an .xlsx) and return it."""
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for tab in spec.tabs:
        _write_tab(wb.create_sheet(_sheet_name(tab.name, used)), tab, theme)
    wb.properties.title = spec.title
    wb.properties.creator = theme.company or "Mavis AI"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path
