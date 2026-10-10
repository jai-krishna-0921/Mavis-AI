"""Draw a DocSpec as a DOCX: styled headings, a cover block, shaded callouts and clean tables.

Built on python-docx's default template with its styles restyled to the theme (fonts, colours, spacing), so
the output stays an ordinary, editable Word document that Google Docs imports faithfully.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

from mavis.studio.spec import Block, DocSpec
from mavis.studio.themes import Theme

CONTENT_WIDTH = 6.5  # inches between 1in margins on Letter; Word reflows on A4


def _hex(color: str) -> str:
    return color.lstrip("#").upper()


def _rgb(color: str) -> RGBColor:
    return RGBColor.from_string(_hex(color))


def _mix(a: str, b: str, t: float) -> str:
    ca, cb = (tuple(int(c.lstrip("#")[i : i + 2], 16) for i in (0, 2, 4)) for c in (a, b))
    return "#" + "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(ca, cb, strict=True))


def _set_font(style, name: str) -> None:
    """Set every script's font and drop theme-font overrides that would win over it."""
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)


def _border_xml(tag: str, color: str, size: int, space: int = 0) -> OxmlElement:
    el = OxmlElement(tag)
    el.set(qn("w:val"), "single" if size else "nil")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), _hex(color))
    return el


def _para_border(paragraph, side: str, color: str, size: int, space: int) -> None:
    """A paragraph border on one side ('left' or 'bottom')."""
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    pbdr.append(_border_xml(f"w:{side}", color, size, space))


def _shade(cell, color: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), _hex(color))
    tcpr.append(shd)


def _cell_borders(cell, **sides: tuple[str, int]) -> None:
    """Borders per side as (colour, eighth-points); unlisted sides get none."""
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        color, size = sides.get(side, ("#FFFFFF", 0))
        borders.append(_border_xml(f"w:{side}", color, size))
    tcpr.append(borders)


def _cell_margins(cell, left: float, right: float, top: float, bottom: float) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, inches in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(inches * 1440)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcpr.append(mar)


def _field(paragraph, instr: str, size: Pt, color: RGBColor, font: str) -> None:
    """Insert a complex field (e.g. PAGE) as three runs."""
    for kind, text in (("begin", ""), (None, instr), ("end", "")):
        run = paragraph.add_run()
        run.font.size, run.font.color.rgb, run.font.name = size, color, font
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


class _Writer:
    def __init__(self, theme: Theme, doc) -> None:
        self.t = theme
        self.doc = doc

    # --- styles ---

    def style_document(self) -> None:
        t, styles = self.t, self.doc.styles
        normal = styles["Normal"]
        _set_font(normal, t.body_font)
        normal.font.size = Pt(11)
        normal.font.color.rgb = _rgb(t.text)
        normal.paragraph_format.line_spacing = 1.15
        normal.paragraph_format.space_after = Pt(6)
        for name, size, before in (("Heading 1", 20, 22), ("Heading 2", 15, 14), ("Title", 34, 0)):
            st = styles[name]
            _set_font(st, t.heading_font)
            st.font.size, st.font.bold, st.font.italic = Pt(size), True, False
            st.font.color.rgb = _rgb(t.primary)
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(8 if name != "Title" else 4)
            st.paragraph_format.line_spacing = 1.05
            ppr = st.element.find(qn("w:pPr"))
            if ppr is not None and ppr.find(qn("w:pBdr")) is not None:
                ppr.remove(ppr.find(qn("w:pBdr")))
        for name in ("List Bullet", "List Number"):
            styles[name].paragraph_format.space_after = Pt(3)
        sec = self.doc.sections[0]
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)

    def footer(self) -> None:
        t = self.t
        sec = self.doc.sections[0]
        style = self.doc.styles["Footer"]
        style.paragraph_format.tab_stops.clear_all()
        p = sec.footer.paragraphs[0]
        p.style = style
        p.paragraph_format.tab_stops.add_tab_stop(Inches(CONTENT_WIDTH), WD_TAB_ALIGNMENT.RIGHT)
        size, color = Pt(9), _rgb(t.muted)
        if t.company:
            r = p.add_run(t.company)
            r.font.size, r.font.color.rgb, r.font.name = size, color, t.body_font
        r = p.add_run("\tPage ")
        r.font.size, r.font.color.rgb, r.font.name = size, color, t.body_font
        _field(p, "PAGE", size, color, t.body_font)

    # --- cover and summary ---

    def cover(self, spec: DocSpec, author: str, date_text: str) -> None:
        t = self.t
        title = self.doc.add_paragraph(spec.title, style="Title")
        title.paragraph_format.space_before = Pt(36)
        if spec.subtitle:
            p = self.doc.add_paragraph()
            r = p.add_run(spec.subtitle)
            r.font.size, r.font.color.rgb, r.font.name = Pt(15), _rgb(t.muted), t.body_font
            p.paragraph_format.space_after = Pt(10)
        meta = " · ".join(x for x in (author, date_text) if x)
        if meta:
            p = self.doc.add_paragraph()
            r = p.add_run(meta)
            r.font.size, r.font.color.rgb = Pt(11), _rgb(t.muted)
            p.paragraph_format.space_after = Pt(4)
        rule = self.doc.add_paragraph()
        rule.paragraph_format.space_after = Pt(18)
        _para_border(rule, "bottom", t.accent, 18, 1)

    def summary(self, text: str) -> None:
        t = self.t
        cell = self._box(t.surface, t.primary)
        label = cell.paragraphs[0]
        r = label.add_run("Key takeaways")
        r.bold = True
        r.font.size, r.font.color.rgb, r.font.name = Pt(12), _rgb(t.primary), t.heading_font
        label.paragraph_format.space_after = Pt(4)
        for line in [x.strip() for x in text.splitlines() if x.strip()]:
            bullet = line[:2] in ("- ", "* ", "• ")
            p = cell.add_paragraph(style="List Bullet" if bullet else None)
            p.add_run(line[2:].strip() if bullet else line)
            p.paragraph_format.space_after = Pt(3)
        self._gap()

    # --- blocks ---

    def _gap(self, pts: int = 6) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run()
        r.font.size = Pt(pts)

    def _box(self, fill: str, bar: str):
        """A one-cell shaded table with a thick left border; returns its cell."""
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Inches(CONTENT_WIDTH)
        _shade(cell, fill)
        _cell_borders(cell, left=(bar, 36))
        _cell_margins(cell, 0.2, 0.18, 0.14, 0.1)
        return cell

    def paragraph(self, text: str) -> None:
        for chunk in [c for c in text.split("\n\n") if c.strip()] or [""]:
            self.doc.add_paragraph(" ".join(chunk.split("\n")).strip())

    def items(self, items: list[str], style: str) -> None:
        paras = [self.doc.add_paragraph(i, style=style) for i in items if i.strip()]
        if style == "List Number" and paras:
            self._restart_numbering(paras)

    def _restart_numbering(self, paras: list[Paragraph]) -> None:
        """Give this numbered list its own counter so it starts at 1."""
        try:
            numbering = self.doc.part.numbering_part.numbering_definitions._numbering
            style_num = paras[0].style.element.pPr.numPr.numId.val
            abstract = numbering.num_having_numId(style_num).abstractNumId.val
            num = numbering.add_num(abstract)
            num.add_lvlOverride(ilvl=0).add_startOverride(1)
            for p in paras:
                numpr = p._p.get_or_add_pPr().get_or_add_numPr()
                numpr.get_or_add_ilvl().val = 0
                numpr.get_or_add_numId().val = num.numId
        except (AttributeError, KeyError):
            pass  # keep the style's shared counter

    def callout(self, block: Block) -> None:
        t = self.t
        cell = self._box(t.surface, t.accent)
        lines = [block.text, *block.items] if block.text else list(block.items)
        for i, line in enumerate(x for x in lines if x.strip()):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.add_run(line)
            p.paragraph_format.space_after = Pt(3)
        self._gap()

    def quote(self, block: Block) -> None:
        t = self.t
        text = block.text or " ".join(block.items)
        p = self.doc.add_paragraph()
        r = p.add_run(text)
        r.italic = True
        r.font.size, r.font.color.rgb = Pt(12), _rgb(t.text)
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.right_indent = Inches(0.3)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(10)
        _para_border(p, "left", t.accent, 24, 12)

    def table(self, block: Block) -> None:
        t = self.t
        header = list(block.table_header)
        ncols = max([len(header), *(len(r) for r in block.table_rows)], default=0)
        if ncols == 0:
            return
        header += [""] * (ncols - len(header))
        rows = [list(r) + [""] * (ncols - len(r)) for r in block.table_rows]
        has_header = any(header)
        grid = ([header] if has_header else []) + rows
        table = self.doc.add_table(rows=len(grid), cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        line = _mix(t.background, t.muted, 0.3)
        for r_i, row in enumerate(grid):
            is_head = has_header and r_i == 0
            tr = table.rows[r_i]._tr
            trpr = tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if is_head:
                trpr.append(OxmlElement("w:tblHeader"))
            band = r_i - (1 if has_header else 0)
            for c_i, text in enumerate(row):
                cell = table.cell(r_i, c_i)
                cell.width = Inches(CONTENT_WIDTH / ncols)
                if is_head:
                    _shade(cell, t.primary)
                elif band % 2 == 0:
                    _shade(cell, t.surface)
                _cell_borders(cell, bottom=(line, 4))
                _cell_margins(cell, 0.1, 0.1, 0.05, 0.05)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                run = p.add_run(text)
                run.font.size = Pt(10.5)
                if is_head:
                    run.bold = True
                    run.font.color.rgb = _rgb("#FFFFFF")
        self._gap(8)

    def block(self, block: Block) -> None:
        kind = block.kind
        if kind == "paragraph":
            self.paragraph(block.text)
            return
        if kind == "quote":
            self.quote(block)
            return
        if kind == "callout":
            self.callout(block)
            return
        if block.text:
            self.paragraph(block.text)
        if kind == "bullets":
            self.items(block.items, "List Bullet")
        elif kind == "numbered":
            self.items(block.items, "List Number")
        elif kind == "table":
            self.table(block)


def render_doc(spec: DocSpec, theme: Theme, path: Path, *, author: str = "", date_text: str = "") -> Path:
    """Write the document to `path` (a .docx) and return it."""
    doc = Document()
    w = _Writer(theme, doc)
    w.style_document()
    w.footer()
    if spec.cover:
        w.cover(spec, author, date_text)
    elif spec.title:
        doc.add_paragraph(spec.title, style="Title")
    if spec.summary.strip():
        w.summary(spec.summary)
    for section in spec.sections:
        if section.heading.strip():
            doc.add_heading(section.heading, level=section.level)
        for block in section.blocks:
            w.block(block)
    doc.core_properties.title = spec.title
    doc.core_properties.author = author or theme.company or "Mavis AI"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
